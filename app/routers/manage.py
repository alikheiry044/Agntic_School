"""روت‌های مدیریت داده (مدیر) — دانش‌آموزان، ورود اکسل، نمرات، غیاب، شهریه.
لیست دانش‌آموزان برای فرم‌های معلم هم باز است؛ تغییرات فقط دست مدیر."""
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from pydantic import BaseModel
from sqlalchemy.orm import Session
from datetime import date
from app.services.telegram_service import ...
from app.database import get_db
from app.models import Student, Grade, Attendance, Tuition, Teacher
# from app.auth import require_admin, require_admin_or_teacher, hash_password
from app.routers.auth import require_admin, require_admin_or_teacher, hash_password
router = APIRouter(prefix="/manage", tags=["Manage"])
# from app.agent.analytics import get_school_wide_report

# ===== Schemas =====
class StudentCreate(BaseModel):
    name: str
    class_name: str
    parent_name: str
    parent_phone: Optional[str] = None
    password: str = "1234"


class StudentUpdate(BaseModel):
    name: str
    class_name: str
    parent_name: str
    parent_phone: Optional[str] = None


class GradeCreate(BaseModel):
    student_id: int
    subject: str
    score: float
    date: str  # YYYY-MM-DD


class AttendanceCreate(BaseModel):
    student_id: int
    date: str
    status: str  # present / absent


class TuitionCreate(BaseModel):
    student_id: int
    amount: float
    due_date: str
    description: Optional[str] = None


# ===== Students =====
@router.get("/students")
def list_students(db: Session = Depends(get_db), user: dict = Depends(require_admin_or_teacher)):
    students = db.query(Student).order_by(Student.id).all()
    return [
        {
            "id": s.id,
            "name": s.name,
            "class_name": s.class_name,
            "parent_name": s.parent_name,
            "parent_phone": getattr(s, "parent_phone", None),
            "has_telegram": bool(getattr(s, "telegram_chat_id", None)),
        }
        for s in students
    ]


@router.post("/students")
def create_student(data: StudentCreate, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = Student(
        name=data.name,
        class_name=data.class_name,
        parent_name=data.parent_name,
        hashed_password=hash_password(data.password),
    )
    if data.parent_phone and hasattr(Student, "parent_phone"):
        student.parent_phone = data.parent_phone
    db.add(student)
    db.commit()
    db.refresh(student)
    return {"message": "دانش‌آموز اضافه شد", "id": student.id}


@router.put("/students/{student_id}")
def update_student(student_id: int, data: StudentUpdate, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="دانش‌آموز پیدا نشد")
    student.name = data.name
    student.class_name = data.class_name
    student.parent_name = data.parent_name
    if data.parent_phone and hasattr(Student, "parent_phone"):
        student.parent_phone = data.parent_phone
    db.commit()
    return {"message": "اطلاعات بروزرسانی شد", "id": student.id}


@router.delete("/students/{student_id}")
def delete_student(student_id: int, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = db.query(Student).filter(Student.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="دانش‌آموز پیدا نشد")
    # حذف داده‌های وابسته تا رکورد یتیم نماند
    db.query(Grade).filter(Grade.student_id == student_id).delete()
    db.query(Attendance).filter(Attendance.student_id == student_id).delete()
    db.query(Tuition).filter(Tuition.student_id == student_id).delete()
    db.delete(student)
    db.commit()
    return {"message": "دانش‌آموز حذف شد", "id": student_id}


# ===== Excel Import =====
@router.post("/students/import")
def import_students(file: UploadFile = File(...), db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    """ورود گروهی از اکسل — ستون‌ها: نام، کلاس، نام والد، رمز (اختیاری)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500, detail="کتابخانه openpyxl نصب نیست — دستور: pip install openpyxl")

    try:
        wb = load_workbook(file.file, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))  # ردیف اول = عنوان
    except Exception:
        raise HTTPException(status_code=400, detail="فایل خوانده نشد — فرمت xlsx معتبر باشد")

    created, errors = 0, []
    for i, row in enumerate(rows, start=2):
        if not row or not row[0]:
            continue
        try:
            name = str(row[0]).strip()
            class_name = str(row[1]).strip()
            parent_name = str(row[2]).strip()
            password = str(row[3]).strip() if len(row) > 3 and row[3] else "1234"
            student = Student(
                name=name,
                class_name=class_name,
                parent_name=parent_name,
                hashed_password=hash_password(password),
            )
            db.add(student)
            db.commit()
            created += 1
        except Exception as e:
            db.rollback()
            errors.append({"row": i, "error": str(e)})

    return {"total_created": created, "errors": errors}


# ===== Grades =====
@router.post("/grades")
def create_grade(data: GradeCreate, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = db.query(Student).filter(Student.id == data.student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="دانش‌آموز پیدا نشد")
    try:
        grade_date = date.fromisoformat(data.date)
    except ValueError:
        raise HTTPException(status_code=400, detail="تاریخ نامعتبر است — فرمت: YYYY-MM-DD")
    if data.score < 0 or data.score > 20:
        raise HTTPException(status_code=400, detail="نمره باید بین ۰ تا ۲۰ باشد")
    grade = Grade(student_id=data.student_id, subject=data.subject, score=data.score, date=grade_date)
    db.add(grade)
    db.commit()
    db.refresh(grade)
    return {"message": "نمره ثبت شد", "id": grade.id}


# ===== Attendance =====
@router.post("/attendance")
def create_attendance(data: AttendanceCreate, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = db.query(Student).filter(Student.id == data.student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="دانش‌آموز پیدا نشد")
    try:
        att_date = date.fromisoformat(data.date)
    except ValueError:
        raise HTTPException(status_code=400, detail="تاریخ نامعتبر است — فرمت: YYYY-MM-DD")
    if data.status not in ("present", "absent"):
        raise HTTPException(status_code=400, detail="وضعیت باید present یا absent باشد")
    att = Attendance(student_id=data.student_id, date=att_date, status=data.status)
    db.add(att)
    db.commit()
    db.refresh(att)
    return {"message": "حضور و غیاب ثبت شد", "id": att.id}


# ===== Tuition =====
@router.post("/tuition")
def create_tuition(data: TuitionCreate, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    student = db.query(Student).filter(Student.id == data.student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="دانش‌آموز پیدا نشد")
    try:
        due = date.fromisoformat(data.due_date)
    except ValueError:
        raise HTTPException(status_code=400, detail="تاریخ نامعتبر است — فرمت: YYYY-MM-DD")
    if data.amount <= 0:
        raise HTTPException(status_code=400, detail="مبلغ باید بزرگ‌تر از صفر باشد")
    tuition = Tuition(
        student_id=data.student_id,
        amount=data.amount,
        due_date=due,
        description=data.description or "قسط شهریه",
        is_paid=False,
    )
    db.add(tuition)
    db.commit()
    db.refresh(tuition)
    return {"message": "قسط ثبت شد", "id": tuition.id}


@router.get("/tuition/student/{student_id}")
def list_tuitions(student_id: int, db: Session = Depends(get_db), user: dict = Depends(require_admin_or_teacher)):
    tuitions = (
        db.query(Tuition)
        .filter(Tuition.student_id == student_id)
        .order_by(Tuition.due_date)
        .all()
    )
    return [
        {
            "id": t.id,
            "amount": t.amount,
            "due_date": t.due_date.isoformat() if t.due_date else None,
            "description": t.description,
            "is_paid": t.is_paid,
        }
        for t in tuitions
    ]


@router.post("/tuition/{tuition_id}/pay")
def pay_tuition(tuition_id: int, db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    tuition = db.query(Tuition).filter(Tuition.id == tuition_id).first()
    if not tuition:
        raise HTTPException(status_code=404, detail="قسط پیدا نشد")
    tuition.is_paid = True
    db.commit()
    return {"message": "پرداخت ثبت شد", "id": tuition.id}

# ===== Teachers =====
from .models import Teacher

@router.get("/teachers")
def list_teachers(db: Session = Depends(get_db), user: dict = Depends(require_admin)):
    teachers = db.query(Teacher).order_by(Teacher.id).all()
    return [
        {
            "id": t.id,
            "name": t.name,
            "subject": t.subject,
            "class_name": t.class_name,
            "username": t.username,
        }
        for t in teachers
    ]


@router.post("/teachers")
def create_teacher(
    name: str = None,
    db: Session = Depends(get_db),
    user: dict = Depends(require_admin),
):
    # این نسخه ساده است — نسخه‌ی کامل پایین‌تر
    ...
